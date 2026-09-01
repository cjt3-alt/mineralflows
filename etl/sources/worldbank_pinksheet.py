"""World Bank Pink Sheet — monthly commodity prices, averaged to a year.

The Pink Sheet is the reference series for exchange-traded metals, so where it
covers one of our minerals it takes precedence over everything else.

Two things about it are worth knowing before reading the code:

* **It covers copper and nothing else we track.** There is no lithium, cobalt or
  rare-earth series in it. Rather than hardcode that, the module matches Pink
  Sheet column headers against mineral names from the config and reports which
  minerals it could not cover. If the World Bank adds lithium tomorrow, this
  module picks it up with no code change.
* **The download URL carries a release hash that changes every month.** So the
  URL is discovered from the commodity markets landing page, with the last
  known-good URL as a fallback. A hardcoded hash would rot within weeks.
"""

from __future__ import annotations

import re
from datetime import date

import openpyxl
import requests

from sources.base import AUTO, ExtractContext, SourceResult, cache_path, skipped

SOURCE_ID = "worldbank-pinksheet"
KIND = AUTO

LANDING_URL = "https://www.worldbank.org/en/research/commodity-markets"
#: Used only if the landing page stops advertising the workbook. Correct as of
#: the 2026-01 release; expected to go stale, which is why it is the fallback.
FALLBACK_XLSX_URL = (
    "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026"
    "/related/CMO-Historical-Data-Monthly.xlsx"
)
LINK_RE = re.compile(
    r"https://thedocs\.worldbank\.org/[^\"'<>\s]*CMO-Historical-Data-Monthly\.xlsx"
)
NAME = "World Bank Commodity Markets Pink Sheet"
SHEET = "Monthly Prices"
#: The only unit this module will accept. Anything else is a silent order-of-
#: magnitude error waiting to happen, so it refuses the column instead.
EXPECTED_UNIT = "$/mt"
USER_AGENT = "Mozilla/5.0 (compatible; mineralflows-etl/1.0)"
TIMEOUT = 180


def discover_url() -> str:
    response = requests.get(LANDING_URL, timeout=TIMEOUT, headers={"User-Agent": USER_AGENT})
    response.raise_for_status()
    match = LINK_RE.search(response.text)
    return match.group(0) if match else FALLBACK_XLSX_URL


def extract(ctx: ExtractContext) -> SourceResult:
    coverage_stub = "Annual averages of the monthly Pink Sheet series, in US dollars per tonne."
    if ctx.offline:
        return skipped(SOURCE_ID, NAME, LANDING_URL, "monthly", coverage_stub, "--offline was set")

    try:
        url = discover_url()
        response = requests.get(url, timeout=TIMEOUT, headers={"User-Agent": USER_AGENT})
        response.raise_for_status()
        path = cache_path(SOURCE_ID, "CMO-Historical-Data-Monthly.xlsx")
        path.write_bytes(response.content)
    except Exception as exc:
        return skipped(
            SOURCE_ID, NAME, LANDING_URL, "monthly", coverage_stub, f"could not download: {exc}"
        )

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            return skipped(
                SOURCE_ID,
                NAME,
                LANDING_URL,
                "monthly",
                coverage_stub,
                f"the workbook has no {SHEET!r} sheet; its layout has changed",
            )
        rows = list(book[SHEET].iter_rows(values_only=True))
    finally:
        book.close()

    # Row 5 holds the series names, row 6 the units, and the monthly observations
    # start at row 7 with a "1960M01" style period label in the first column.
    headers = [str(c).strip() if c is not None else "" for c in rows[4]]
    units = [str(c).strip() if c is not None else "" for c in rows[5]]

    wanted = {m["name"].strip().lower(): m["id"] for m in ctx.minerals if m.get("active")}
    columns: dict[str, int] = {}
    warnings: list[str] = []

    for index, header in enumerate(headers):
        mineral_id = wanted.get(header.lower())
        if mineral_id is None:
            continue
        unit = units[index].strip("() ")
        if unit != EXPECTED_UNIT:
            warnings.append(
                f"{header} is quoted in {unit!r}, not {EXPECTED_UNIT!r}; the column was ignored "
                "rather than converted on an assumption"
            )
            continue
        columns[mineral_id] = index

    uncovered = sorted(set(wanted.values()) - set(columns))
    if uncovered:
        warnings.append(
            "no Pink Sheet series for " + ", ".join(uncovered) + "; those prices come from "
            "another source or from the seed"
        )

    monthly: dict[tuple[str, int], list[float]] = {}
    for row in rows[6:]:
        label = row[0]
        if not label:
            continue
        match = re.match(r"^(\d{4})M(\d{2})$", str(label).strip())
        if not match:
            continue
        year = int(match.group(1))
        for mineral_id, index in columns.items():
            value = row[index]
            try:
                number = float(value)
            except (TypeError, ValueError):
                continue  # the sheet writes an ellipsis for missing months
            monthly.setdefault((mineral_id, year), []).append(number)

    years = [ctx.year - 1, ctx.year]
    prices: list[dict] = []
    for (mineral_id, year), values in sorted(monthly.items()):
        if year not in years:
            continue
        if len(values) < 12:
            warnings.append(
                f"{mineral_id} {year} has only {len(values)} monthly observations; "
                "a partial year is not an annual average, so it was dropped"
            )
            continue
        prices.append(
            {
                "mineral_id": mineral_id,
                "year": year,
                "avg_price_usd_per_tonne": round(sum(values) / len(values), 2),
                "source": SOURCE_ID,
                "source_url": LANDING_URL,
            }
        )

    covered = sorted({p["mineral_id"] for p in prices})
    return SourceResult(
        source_id=SOURCE_ID,
        name=NAME,
        url=LANDING_URL,
        vintage=f"{min(years)}-{max(years)} annual averages",
        retrieved_at=date.today(),
        coverage=(
            coverage_stub
            + " Covers "
            + (", ".join(covered) if covered else "none of the tracked minerals")
            + "."
        ),
        prices=prices,
        warnings=warnings,
    )
