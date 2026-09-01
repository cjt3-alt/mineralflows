"""IEA Critical Minerals Data Explorer — the mine-versus-refine split.

A manual drop: the IEA serves the workbook from a data tool that blocks
automated fetches, so a human downloads it into `etl/raw/manual/iea/`. Licensed
CC BY 4.0, which is why the copy here can be committed.

What it adds that nothing else does: refining capacity by country for lithium and
cobalt. The USGS publishes mine production for all four minerals but refinery
production only for copper, so without this there is no way to see that the DRC
mines the cobalt and China refines it.

Like the USGS production figures, this lands nowhere in `public/data/` — the
contract has no production table and adding one is out of scope. It feeds the
cross-source check that flags a trade flow larger than the origin country's
entire output, which is the check most likely to catch a unit error.

The workbook publishes projection years starting at 2025 rather than historical
years, so an exact match for the trade year is usually unavailable. The nearest
published year is used and named in the coverage string, because the alternative
is a check that never runs.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

from contract import ETL_ROOT
from sources.base import (
    MANUAL,
    ExtractContext,
    SourceResult,
    drop_date,
    find_manual,
    manual_dir,
    skipped,
)

SOURCE_ID = "iea-critical-minerals"
KIND = MANUAL

NAME = "IEA Critical Minerals Data Explorer"
URL = "https://www.iea.org/data-and-statistics/data-tools/critical-minerals-data-explorer"
EXPECTED_FILENAME = "*Critical Minerals*.xlsx"
SHEET_PREFIX = "2 Total supply"
ALIAS_FILE = ETL_ROOT / "config" / "commodity_aliases.json"

#: The sheet lays two tables side by side, headed "<Mineral> - Mining" and
#: "<Mineral> - Refining", with a blank spacer column between them.
LABEL_RE = re.compile(r"^(?P<mineral>.+?)\s*-\s*(?P<stage>Mining|Refining)\b", re.IGNORECASE)
STAGE_WORDS = {"mining": "mine", "refining": "refine"}
UNIT_RE = re.compile(r"\((?P<unit>[a-zA-Z]+)\)")
UNIT_TO_TONNES = {"kt": 1000.0, "mt": 1_000_000.0, "t": 1.0}

#: Roll-up rows inside each table. They are not countries.
NON_COUNTRY_ROWS = {"rest of world", "total", "top 3 share", "total clean technologies"}


def extract(ctx: ExtractContext) -> SourceResult:
    coverage_stub = (
        "Country-level mining and refining supply, used to sanity-check trade volumes "
        "against production. Not written to any output file."
    )
    if ctx.ignore_manual:
        return skipped(SOURCE_ID, NAME, URL, "unknown", coverage_stub, "--no-manual was set")

    matches = find_manual(SOURCE_ID, EXPECTED_FILENAME)
    if not matches:
        return skipped(
            SOURCE_ID,
            NAME,
            URL,
            "unknown",
            coverage_stub,
            f"no file matching {EXPECTED_FILENAME} in {manual_dir(SOURCE_ID)}. "
            "See etl/README.md for where to download it.",
        )

    path = matches[-1]
    try:
        return _read(path, ctx, coverage_stub)
    except Exception as exc:
        return skipped(
            SOURCE_ID, NAME, URL, "unknown", coverage_stub, f"could not read {path.name}: {exc}"
        )


def _read(path: Path, ctx: ExtractContext, coverage_stub: str) -> SourceResult:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = next((s for s in book.sheetnames if s.startswith(SHEET_PREFIX)), None)
        if sheet_name is None:
            raise RuntimeError(f"no sheet starting with {SHEET_PREFIX!r}")
        rows = [list(r) for r in book[sheet_name].iter_rows(values_only=True)]
    finally:
        book.close()

    unit_factor, unit_text = _unit(rows)
    blocks = _year_blocks(rows)
    if not blocks:
        raise RuntimeError("no year header row was found")
    published = sorted({year for block in blocks for year in block})
    used_year = min(published, key=lambda y: (abs(y - ctx.year), y))

    aliases = json.loads(ALIAS_FILE.read_text(encoding="utf-8"))["tokens"]
    by_label: dict[str, str] = {}
    for mineral in ctx.minerals:
        if not mineral.get("active"):
            continue
        for token in {mineral["name"], *aliases.get(mineral["id"], [])}:
            by_label[token.strip().lower()] = mineral["id"]

    production: dict[tuple[str, str, str], float] = {}
    warnings: list[str] = []
    unresolved = 0
    #: Column of the label that started the table this column block belongs to.
    active: dict[int, tuple[str, str]] = {}

    for row in rows:
        for column, cell in enumerate(row):
            if cell is None:
                continue
            match = LABEL_RE.match(str(cell).strip())
            if not match:
                continue
            mineral_id = by_label.get(match.group("mineral").strip().lower())
            stage = STAGE_WORDS[match.group("stage").lower()]
            if mineral_id is None:
                active.pop(column, None)
            else:
                active[column] = (mineral_id, stage)

        for column, (mineral_id, stage) in list(active.items()):
            label = row[column] if column < len(row) else None
            if label is None or LABEL_RE.match(str(label).strip()):
                continue
            name = str(label).strip()
            if name.lower() in NON_COUNTRY_ROWS or ctx.resolver.is_aggregate(name):
                continue
            value_column = _value_column(blocks, label_column=column, year=used_year)
            if value_column is None or value_column >= len(row):
                continue
            value = row[value_column]
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                continue
            iso3 = ctx.resolver.resolve(name)
            if iso3 is None:
                unresolved += 1
                continue
            key = (mineral_id, iso3, stage)
            production[key] = production.get(key, 0.0) + value * unit_factor

    if unresolved:
        warnings.append(f"{unresolved} supply rows named a country that did not resolve")
    if used_year != ctx.year:
        warnings.append(
            f"the workbook publishes {published}, so {used_year} was used in place of "
            f"{ctx.year} for the production cross-check"
        )

    return SourceResult(
        source_id=SOURCE_ID,
        name=NAME,
        url=URL,
        vintage=str(used_year),
        retrieved_at=drop_date(path),
        coverage=(
            f"{len(production)} country-stage figures for {used_year}, read in {unit_text}. "
            + coverage_stub
        ),
        production=production,
        warnings=warnings,
    )


def _unit(rows: list[list]) -> tuple[float, str]:
    """Read the unit out of the sheet title rather than assuming kilotonnes."""
    for row in rows[:6]:
        for cell in row:
            if cell is None:
                continue
            match = UNIT_RE.search(str(cell))
            if match:
                unit = match.group("unit").lower()
                if unit in UNIT_TO_TONNES:
                    return UNIT_TO_TONNES[unit], unit
    raise RuntimeError("the sheet title does not state a unit, so the figures cannot be scaled")


def _year_blocks(rows: list[list]) -> list[dict[int, int]]:
    """The year header row, split into one block per side-by-side table.

    Mining and refining share a single header row, so 2025 appears twice in it at
    different columns. Splitting on the blank spacer column keeps each table's
    year columns with the table they belong to.
    """
    for row in rows[:12]:
        found = [
            (column, int(str(cell).strip()))
            for column, cell in enumerate(row)
            if cell is not None and re.fullmatch(r"\d{4}", str(cell).strip())
        ]
        if len(found) < 2:
            continue
        blocks: list[dict[int, int]] = []
        current: dict[int, int] = {}
        previous: int | None = None
        for column, year in found:
            if previous is not None and column != previous + 1:
                blocks.append(current)
                current = {}
            current[year] = column
            previous = column
        blocks.append(current)
        return blocks
    return []


def _value_column(blocks: list[dict[int, int]], label_column: int, year: int) -> int | None:
    """The column holding `year` for the table whose labels sit at `label_column`."""
    candidates = [
        block[year] for block in blocks if year in block and min(block.values()) > label_column
    ]
    return min(candidates) if candidates else None
